"""
Adaptadores de salida de la ingesta y de la IA, sin red: archivos reales en memoria y HTTP simulado.
"""

from __future__ import annotations

import io
import json
from typing import Any

import httpx
import pytest

from platform5e.application.services.generation.prompt_drafting import (
    build_generation_context,
    normalize_generated_resources,
)
from platform5e.application.services.generation.rule_based_generator import RuleBasedGenerator
from platform5e.domain.materials.chunking import MaterialError
from platform5e.infrastructure.extraction.document_text_extractor import DocumentTextExtractor
from platform5e.infrastructure.llm.gemini_language_model import AIProviderError, GeminiLanguageModel
from platform5e.infrastructure.llm.language_model_factory import create_generators, describe_generator
from platform5e.infrastructure.sources.wikipedia_source import WikipediaSource

SAMPLE_CORPUS = (
    "Una red de computadoras es un conjunto de equipos conectados que comparten recursos e información. "
    "El modelo OSI es un modelo de referencia que divide la comunicación en siete capas. "
    "El modelo TCP/IP es el modelo que se usa en internet y tiene cuatro capas. "
    "El protocolo TCP es orientado a la conexión y garantiza una entrega confiable de los datos. "
    "El protocolo UDP no establece conexión y por eso es más rápido para videollamadas. "
    "El enrutador es el dispositivo que decide por qué camino viaja cada paquete hasta su destino. "
    "La encapsulación es el proceso por el cual cada capa agrega su propio encabezado a los datos."
)


def create_minimal_pdf_bytes(text_content: str) -> bytes:
    """PDF de una página con texto (sin tildes, fuente Helvetica)."""
    stream_content = f"BT /F1 12 Tf 72 720 Td ({text_content}) Tj ET".encode()
    pdf_objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream_content) + stream_content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    buffer, offsets = io.BytesIO(), []
    buffer.write(b"%PDF-1.4\n")
    for idx, obj in enumerate(pdf_objects, start=1):
        offsets.append(buffer.tell())
        buffer.write(b"%d 0 obj\n" % idx + obj + b"\nendobj\n")
    xref_pos = buffer.tell()
    buffer.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(pdf_objects) + 1))
    buffer.writelines(b"%010d 00000 n \n" % p for p in offsets)
    buffer.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF" % (len(pdf_objects) + 1, xref_pos))
    return buffer.getvalue()


def create_minimal_pptx_bytes() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Protocolos"
    slide.placeholders[1].text = "TCP garantiza la entrega confiable de los datos en la red."
    slide.shapes.add_table(1, 2, Inches(1), Inches(4), Inches(4), Inches(1)).table.cell(0, 0).text = "Capa de transporte"
    slide.notes_slide.notes_text_frame.text = "Explicar con un ejemplo."
    prs.slides.add_slide(prs.slide_layouts[6])
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def test_extractor_pdf_pptx_y_errores():
    extractor = DocumentTextExtractor()
    assert extractor.extract("pdf", create_minimal_pdf_bytes("El protocolo TCP garantiza la entrega.")) == [
        ("p. 1", "El protocolo TCP garantiza la entrega.")
    ]
    pages = extractor.extract("pptx", create_minimal_pptx_bytes())
    assert pages[0][0] == "diapositiva 1" and "TCP garantiza" in pages[0][1]
    assert "Capa de transporte" in pages[0][1] and "Notas: Explicar" in pages[0][1] and pages[1] == ("diapositiva 2", "")

    error_cases = [
        ("pdf", b"%PDF-1.4 roto", "ARCHIVO_ILEGIBLE"),
        ("pptx", b"no es zip", "ARCHIVO_ILEGIBLE"),
        ("txt", b"\xff\xfe\x00", "FORMATO_INVALIDO"),
        ("doc", b"x", "FORMATO_INVALIDO"),
    ]
    for file_type, content, expected_code in error_cases:
        with pytest.raises(MaterialError) as err:
            extractor.extract(file_type, content)
        assert err.value.code == expected_code


def test_extractor_txt_por_secciones_y_por_titulos_de_wikipedia():
    extractor = DocumentTextExtractor()
    long_body = "\n\n".join(["Párrafo de prueba con varias palabras. " * 20] * 6)
    sections = extractor.extract("txt", long_body.encode("utf-8-sig"))
    assert len(sections) > 1 and sections[0][0] == "sección 1"

    wiki_content = "Introducción del artículo.\n\n== Historia ==\nOrigen de la red.\n\n=== Detalle ===\nMás datos."
    wiki_sections = [(loc, text.strip()) for loc, text in extractor.extract("txt", wiki_content.encode())]
    assert wiki_sections == [
        ("introducción", "Introducción del artículo."),
        ("sección «Historia»", "Origen de la red."),
        ("sección «Detalle»", "Más datos."),
    ]


def build_mock_http_client(responses_list: list[tuple[int, dict[str, Any]]]) -> tuple[httpx.Client, list[httpx.Request]]:
    captured_requests: list[httpx.Request] = []

    def mock_handler(request: httpx.Request) -> httpx.Response:
        captured_requests.append(request)
        status_code, body = responses_list[min(len(captured_requests), len(responses_list)) - 1]
        return httpx.Response(status_code, json=body)

    return httpx.Client(transport=httpx.MockTransport(mock_handler)), captured_requests


TEST_GENERATION_CONTEXT = build_generation_context(
    {"name": "Redes"},
    {"id": "u1", "number": 1, "title": "Protocolos", "outcomes": []},
    {"stage": "evaluate", "resourceType": "item_opcion_multiple", "quantity": 1, "optionCount": 4},
    [{"id": "f1", "documento_id": "d", "documentId": "d", "ubicacion": "p. 1", "location": "p. 1", "texto": SAMPLE_CORPUS, "text": SAMPLE_CORPUS}],
    {"d": "apuntes.txt"},
)
EXPECTED_LLM_OUTPUT = {
    "recursos": [{"titulo": "T", "contenido": "C", "citas": [{"afirmacion": "a", "fragmentos": ["F1"]}], "alternativas": []}]
}


def test_gemini_envia_instrucciones_y_lee_el_texto():
    client, requests = build_mock_http_client(
        [(200, {"candidates": [{"content": {"parts": [{"text": json.dumps(EXPECTED_LLM_OUTPUT)}]}}]})]
    )
    generator = GeminiLanguageModel("gemini", "clave", client=client)
    output = generator.generate(TEST_GENERATION_CONTEXT)
    assert output == EXPECTED_LLM_OUTPUT
    assert requests[0].url.path.endswith("/models/gemini-2.5-flash:generateContent") or requests[0].url.path.endswith("/generateContent")
    assert requests[0].headers["x-goog-api-key"] == "clave"
    payload = json.loads(requests[0].content)
    assert payload["generationConfig"]["responseMimeType"] == "application/json"
    assert "[F1]" in payload["contents"][0]["parts"][0]["text"]


def test_gemini_pide_salida_estructurada_y_sigue_sin_ella_si_el_modelo_la_rechaza():
    json_text = json.dumps(EXPECTED_LLM_OUTPUT)
    client, requests = build_mock_http_client(
        [
            (400, {"error": {"message": "Invalid JSON payload: response_schema"}}),
            (200, {"candidates": [{"content": {"parts": [{"text": json_text}]}}]}),
        ]
    )
    generator = GeminiLanguageModel("gemini", "clave", client=client)
    assert generator.generate(TEST_GENERATION_CONTEXT) == EXPECTED_LLM_OUTPUT
    first_config, second_config = (json.loads(r.content)["generationConfig"] for r in requests)
    assert first_config["responseSchema"]["required"] == ["recursos"] and "responseSchema" not in second_config
    assert generator.generate(TEST_GENERATION_CONTEXT) == EXPECTED_LLM_OUTPUT and len(requests) == 3


def test_anthropic_y_openai_compatible():
    client_anthropic, requests_anthropic = build_mock_http_client(
        [(200, {"content": [{"type": "text", "text": json.dumps(EXPECTED_LLM_OUTPUT)}]})]
    )
    gen_anthropic = GeminiLanguageModel("anthropic", "clave", client=client_anthropic)
    assert gen_anthropic.generate(TEST_GENERATION_CONTEXT) == EXPECTED_LLM_OUTPUT
    assert gen_anthropic.description == "Anthropic · claude-sonnet-5-5"
    assert requests_anthropic[0].url.path == "/v1/messages"
    assert requests_anthropic[0].headers["anthropic-version"] == "2023-06-01"

    client_openai, requests_openai = build_mock_http_client(
        [(200, {"choices": [{"message": {"content": json.dumps(EXPECTED_LLM_OUTPUT["recursos"])}}]})]
    )
    gen_openai = GeminiLanguageModel(
        "openai", "clave", model="llama-3.3", base_url="https://api.groq.com/openai/v1/", client=client_openai
    )
    assert gen_openai.generate(TEST_GENERATION_CONTEXT) == EXPECTED_LLM_OUTPUT
    assert str(requests_openai[0].url) == "https://api.groq.com/openai/v1/chat/completions"
    assert requests_openai[0].headers["authorization"] == "Bearer clave"


def test_reintenta_una_vez_si_el_json_no_es_valido():
    client, requests = build_mock_http_client(
        [
            (200, {"content": [{"type": "text", "text": "no es json"}]}),
            (200, {"content": [{"type": "text", "text": json.dumps(EXPECTED_LLM_OUTPUT)}]}),
        ]
    )
    assert GeminiLanguageModel("anthropic", "k", client=client).generate(TEST_GENERATION_CONTEXT) == EXPECTED_LLM_OUTPUT
    assert len(requests) == 2 and "no era JSON válido" in json.loads(requests[1].content)["messages"][0]["content"]


@pytest.mark.parametrize(
    "status_code,body,expected_substring",
    [
        (429, {"error": {"message": "quota"}}, "límite de uso"),
        (400, {"error": {"message": "API key not valid"}}, "API key not valid"),
        (200, {"promptFeedback": {"blockReason": "SAFETY"}}, "SAFETY"),
    ],
)
def test_errores_del_proveedor_se_explican(status_code: int, body: dict[str, Any], expected_substring: str):
    client, _ = build_mock_http_client([(status_code, body)])
    with pytest.raises(AIProviderError) as err:
        GeminiLanguageModel("gemini", "k", client=client).generate(TEST_GENERATION_CONTEXT)
    assert expected_substring in str(err.value)


def test_sin_conexion_y_proveedor_desconocido():
    def mock_failing_transport(req):
        raise httpx.ConnectError("sin red")

    with pytest.raises(AIProviderError):
        GeminiLanguageModel("openai", "k", client=httpx.Client(transport=httpx.MockTransport(mock_failing_transport))).generate(
            TEST_GENERATION_CONTEXT
        )

    with pytest.raises(ValueError):
        GeminiLanguageModel("otro", "k")


def test_configuracion_desde_el_entorno():
    primary, fallback = create_generators({})
    assert primary is fallback and primary.uses_ai is False
    assert describe_generator(primary) == {
        "descripcion": "Generador por reglas (sin IA)",
        "usaIA": False,
        "proveedor": "reglas",
        "modelo": None,
    }

    primary_with_key, fallback_with_key = create_generators({"IA_API_KEY": "k"})
    assert primary_with_key.provider == "gemini" and fallback_with_key.uses_ai is False

    primary_anthropic, _ = create_generators(
        {"IA_API_KEY": "k", "IA_PROVEEDOR": "Anthropic", "IA_MODELO": "claude-haiku-5-5", "IA_TIEMPO_MAX": "30"}
    )
    assert describe_generator(primary_anthropic)["modelo"] == "claude-haiku-5-5"
    assert create_generators({"IA_API_KEY": "k", "IA_PROVEEDOR": "reglas"})[0].uses_ai is False


def test_wikipedia_busca_y_limpia_el_articulo():
    def mock_wiki_handler(request: httpx.Request) -> httpx.Response:
        if request.url.params.get("list") == "search":
            return httpx.Response(
                200,
                json={"query": {"search": [{"title": "Protocolo TCP"}, {"title": "Desambiguación"}, {"title": "Falta"}]}},
            )
        title = request.url.params["titles"]
        if title == "Falta":
            return httpx.Response(200, json={"query": {"pages": [{"missing": True}]}})
        raw_extract = "Corto." if title == "Desambiguación" else SAMPLE_CORPUS + "\n\n== Referencias ==\nLibro."
        return httpx.Response(
            200,
            json={"query": {"pages": [{"title": title, "extract": raw_extract, "fullurl": "https://es.wikipedia.org/wiki/TCP"}]}},
        )

    source = WikipediaSource(client=httpx.Client(transport=httpx.MockTransport(mock_wiki_handler)))
    articles = source.search("protocolo TCP", 3)
    assert [a["titulo"] for a in articles] == ["Protocolo TCP"]
    assert articles[0]["licencia"] == "CC BY-SA 4.0" and "Referencias" not in articles[0]["texto"]


@pytest.mark.parametrize(
    "stage,resource_type",
    [
        ("evaluate", "item_opcion_multiple"),
        ("engage", "pregunta_detonante"),
        ("engage", "sondeo_diagnostico"),
        ("explore", "guia_exploracion"),
        ("explore", "caso_indagacion"),
        ("explain", "explicacion"),
        ("explain", "glosario"),
        ("elaborate", "ejercicio_aplicacion"),
    ],
)
def test_generador_por_reglas_produce_recursos_validos_y_citados(stage: str, resource_type: str):
    context = build_generation_context(
        {"name": "Redes"},
        {"id": "u1", "number": 1, "title": "Protocolos", "outcomes": []},
        {"stage": stage, "resourceType": resource_type, "quantity": 2, "optionCount": 4},
        [{"id": "f1", "documento_id": "d", "documentId": "d", "ubicacion": "p. 1", "location": "p. 1", "texto": SAMPLE_CORPUS, "text": SAMPLE_CORPUS}],
        {},
    )
    resources = normalize_generated_resources(RuleBasedGenerator().generate(context), context)
    assert resources, resource_type
    assert all(r["citas"][0]["fragmentIds"] == ["f1"] for r in resources)
    if resource_type == "item_opcion_multiple":
        item = resources[0]
        assert len(item["alternativas"]) == 4 and sum(a["correcta"] for a in item["alternativas"]) == 1
        assert "________" in item["contenido"]
        key_text = next(a["texto"] for a in item["alternativas"] if a["correcta"])
        assert key_text.lower() not in item["titulo"].lower()


def test_generador_por_reglas_sin_material_util_no_inventa():
    context = build_generation_context(
        {},
        {"id": "u1", "number": 1, "title": "U", "outcomes": []},
        {"stage": "evaluate", "resourceType": "item_opcion_multiple", "quantity": 1, "optionCount": 4},
        [{"id": "f1", "documento_id": "d", "documentId": "d", "ubicacion": "p. 1", "location": "p. 1", "texto": "Hola.", "text": "Hola."}],
        {},
    )
    assert RuleBasedGenerator().generate(context) == {"recursos": []}