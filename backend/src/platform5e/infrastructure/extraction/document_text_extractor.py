"""
Extractor de texto para archivos PDF, PPTX y TXT (EN-007).
Retorna una colección de tuplas (ubicación, texto) por página, diapositiva o sección.
"""

from __future__ import annotations

import io
import re
from typing import Any

from platform5e.application.ports.outbound.chunk_repository_port import TextExtractorPort
from platform5e.domain.materials.chunking import MaterialError

MAX_PROCESSED_PAGES: int = 400
CHARS_PER_TEXT_SECTION: int = 3000
_WIKI_HEADING_REGEX: re.Pattern[str] = re.compile(r"^\s*(={2,})\s*(.+?)\s*\1\s*$", re.MULTILINE)


class DocumentTextExtractor(TextExtractorPort):
    """Adaptador de infraestructura para extraer texto legible de formatos binarios y planos."""

    def extract(self, file_type: str, content: bytes) -> list[tuple[str, str]]:
        """Extrae el contenido textual estructurado según el tipo de archivo."""
        if file_type == "pdf":
            return self._extract_pdf(content)
        if file_type == "pptx":
            return self._extract_pptx(content)
        if file_type == "txt":
            return self._extract_txt(content)
        raise MaterialError("FORMATO_INVALIDO", "Usa PDF, PPTX o TXT.")

    @staticmethod
    def _extract_pdf(content: bytes) -> list[tuple[str, str]]:
        """Extrae texto de documentos PDF página por página."""
        from pypdf import PdfReader

        try:
            reader = PdfReader(io.BytesIO(content))
            if reader.is_encrypted and not reader.decrypt(""):
                raise MaterialError(
                    "PDF_PROTEGIDO",
                    "El PDF está protegido con contraseña. Quita la protección y vuelve a subirlo.",
                )
            return [
                (f"p. {i}", page.extract_text() or "")
                for i, page in enumerate(reader.pages[:MAX_PROCESSED_PAGES], start=1)
            ]
        except MaterialError:
            raise
        except Exception as exc:
            raise MaterialError(
                "ARCHIVO_ILEGIBLE",
                "No se pudo leer el PDF. Puede estar dañado; ábrelo, guárdalo de nuevo como PDF y vuelve a subirlo.",
            ) from exc

    @staticmethod
    def _extract_pptx(content: bytes) -> list[tuple[str, str]]:
        """Extrae texto de presentaciones PPTX diapositiva por diapositiva."""
        from pptx import Presentation

        try:
            presentation = Presentation(io.BytesIO(content))
        except Exception as exc:
            raise MaterialError(
                "ARCHIVO_ILEGIBLE",
                "No se pudo leer la presentación. Guárdala de nuevo como PPTX y vuelve a subirla.",
            ) from exc

        extracted_pages: list[tuple[str, str]] = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            if slide_number > MAX_PROCESSED_PAGES:
                break
            shape_texts = [text for shape in slide.shapes for text in _extract_shape_text(shape)]
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    shape_texts.append(f"Notas: {notes}")
            extracted_pages.append(
                (
                    f"diapositiva {slide_number}",
                    "\n\n".join(text.strip() for text in shape_texts if text.strip()),
                )
            )
        return extracted_pages

    @staticmethod
    def _extract_txt(content: bytes) -> list[tuple[str, str]]:
        """Extrae y segmenta texto plano considerando encabezados tipo Wiki o bloques de longitud."""
        try:
            decoded_text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise MaterialError("FORMATO_INVALIDO", "Guarda el TXT con codificación UTF-8.") from exc

        wiki_headings = list(_WIKI_HEADING_REGEX.finditer(decoded_text))
        if wiki_headings:
            sections: list[tuple[str, str]] = [("introducción", decoded_text[: wiki_headings[0].start()])]
            for current_match, next_match in zip(wiki_headings, [*wiki_headings[1:], None]):
                end_index = next_match.start() if next_match else len(decoded_text)
                sections.append(
                    (
                        f"sección «{current_match.group(2)}»",
                        decoded_text[current_match.end() : end_index],
                    )
                )
            return [(loc, txt) for loc, txt in sections if txt.strip()][:MAX_PAGINAS_RESTRICCION]

        sections_list: list[str] = []
        current_buffer = ""
        for paragraph in re.split(r"\n\s*\n", decoded_text):
            if current_buffer and len(current_buffer) + len(paragraph) > CHARS_PER_TEXT_SECTION:
                sections_list.append(current_buffer)
                current_buffer = ""
            current_buffer = f"{current_buffer}\n\n{paragraph}" if current_buffer else paragraph

        if current_buffer.strip():
            sections_list.append(current_buffer)

        return [
            (f"sección {idx}", section_text)
            for idx, section_text in enumerate(sections_list[:MAX_PROCESSED_PAGES], start=1)
        ]


MAX_PAGINAS_RESTRICCION = MAX_PROCESSED_PAGES


def _extract_shape_text(shape: Any) -> list[str]:
    """Extrae el contenido textual de formas, tablas y grupos de PowerPoint."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
        return [text for inner_shape in shape.shapes for text in _extract_shape_text(inner_shape)]
    if getattr(shape, "has_table", False) and shape.has_table:
        return [" | ".join(cell.text for cell in row.cells) for row in shape.table.rows]
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        return [shape.text_frame.text]
    return []
    